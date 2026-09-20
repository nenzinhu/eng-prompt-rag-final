"""Extração de texto de documentos e geração de DOCX/PDF (sem I/O de rede)."""
from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader

# ── Extração de texto ───────────────────────────────────────────────────────


def _extrair_paginas_docx(dados: bytes) -> list[tuple[int, str]]:
    """Um .docx não tem conceito nativo de página — cada heading de nível 1 vira
    uma "página" (mantém granularidade parecida com PDF para a citação de fonte).
    Sem heading nenhum, o documento inteiro vira a página 1.
    """
    documento = DocxDocument(io.BytesIO(dados))
    paginas: list[tuple[int, list[str]]] = [(1, [])]
    # Percorre o CORPO na ordem do documento. `documento.paragraphs` pula as
    # tabelas, e juntá-las no fim jogava toda tabela para a última página: a
    # citação apontava o capítulo errado, que é pior do que não indexar.
    for bloco in _blocos_do_corpo(documento):
        if isinstance(bloco, Table):
            paginas[-1][1].extend(_linhas_da_tabela(bloco))
            continue
        texto = bloco.text.strip()
        if not texto:
            continue
        if bloco.style.name == "Heading 1" and paginas[-1][1]:
            paginas.append((paginas[-1][0] + 1, []))
        paginas[-1][1].append(texto)
    return [(num, "\n".join(linhas)) for num, linhas in paginas]


def _blocos_do_corpo(documento):
    """Parágrafos e tabelas do .docx, na ordem em que aparecem no documento."""
    corpo = documento.element.body
    for filho in corpo.iterchildren():
        if filho.tag == qn("w:p"):
            yield Paragraph(filho, documento)
        elif filho.tag == qn("w:tbl"):
            yield Table(filho, documento)


def _linhas_da_tabela(tabela) -> list[str]:
    """Uma linha de texto por linha da tabela, células separadas por " | "."""
    linhas: list[str] = []
    for linha in tabela.rows:
        celulas = [c.text.strip().replace("\n", " ") for c in linha.cells]
        # Célula mesclada aparece repetida nas colunas que ocupa.
        unicas: list[str] = []
        for celula in celulas:
            if celula and (not unicas or unicas[-1] != celula):
                unicas.append(celula)
        if unicas:
            linhas.append(" | ".join(unicas))
    return linhas


def _extrair_paginas_pptx(dados: bytes) -> list[tuple[int, str]]:
    """Cada slide vira uma página — mapeamento natural, igual PDF."""
    apresentacao = Presentation(io.BytesIO(dados))
    paginas = []
    for i, slide in enumerate(apresentacao.slides):
        textos: list[str] = []
        for forma in slide.shapes:
            textos.extend(_texto_da_forma(forma))
        # As notas do apresentador costumam ter a explicação que o slide resume.
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notas = slide.notes_slide.notes_text_frame.text.strip()
            if notas:
                textos.append(notas)
        paginas.append((i + 1, "\n".join(textos)))
    return paginas


def _texto_da_forma(forma) -> list[str]:
    """Texto de uma forma do slide, entrando em grupos e tabelas.

    Antes só se lia `forma.text_frame`, e tabela e grupo ficavam fora do índice,
    embora seja onde o material costuma pôr comparações e resumos. (SmartArt
    continua de fora: o python-pptx não expõe o texto do diagrama.)

    Uma forma só nunca pode derrubar a ingestão do arquivo inteiro, então cada
    ramo é protegido: o `shape_type` do python-pptx levanta
    `NotImplementedError` em forma não reconhecida — e `getattr` com padrão NÃO
    segura isso, porque a exceção vem de dentro da property.
    """
    textos: list[str] = []
    try:
        # Grupo: as formas de dentro é que têm o texto. `shapes` só existe em
        # grupo, então serve de teste sem tocar em `shape_type`.
        if hasattr(forma, "shapes"):
            for interna in forma.shapes:
                textos.extend(_texto_da_forma(interna))
            return textos
        if getattr(forma, "has_table", False):
            for linha in forma.table.rows:
                celulas = [c.text.strip().replace("\n", " ") for c in linha.cells if c.text.strip()]
                if celulas:
                    textos.append(" | ".join(celulas))
            return textos
        if forma.has_text_frame and forma.text_frame.text.strip():
            textos.append(forma.text_frame.text)
    except Exception:
        # Perder uma forma exótica é aceitável; perder o arquivo, não.
        return textos
    return textos


def _texto_gravavel(texto: str) -> str:
    """Tira do texto o que o Postgres recusa numa coluna text.

    O caractere nulo (U+0000) vem de alguns geradores de PDF e derrubava a
    gravação com 22P05 (caso real em 18/09/2026). Surrogate UTF-16 solto, de PDF
    mal formado, também não vira UTF-8 válido e passa a ser trocado por "?".
    """
    return texto.replace("\x00", "").encode("utf-8", "replace").decode("utf-8")


def extrair_paginas(pdf_bytes: bytes, nome: str = "") -> list[tuple[int, str]]:
    """Devolve [(numero_da_pagina, texto)]. Aceita PDF, TXT, MD, DOCX e PPTX."""
    return [(n, _texto_gravavel(t)) for n, t in _extrair_paginas_bruto(pdf_bytes, nome)]


def _extrair_paginas_bruto(pdf_bytes: bytes, nome: str = "") -> list[tuple[int, str]]:
    nome_lower = nome.lower()
    if nome_lower.endswith((".txt", ".md")):
        try:
            texto = pdf_bytes.decode("utf-8")
        except UnicodeDecodeError:
            texto = pdf_bytes.decode("latin-1", "replace")
        return [(1, texto)]
    if nome_lower.endswith(".docx"):
        return _extrair_paginas_docx(pdf_bytes)
    if nome_lower.endswith(".pptx"):
        return _extrair_paginas_pptx(pdf_bytes)
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def montar_contexto(chunks: list[dict]) -> tuple[str, float]:
    contexto = "\n\n".join(f"[Fonte: {c['source']}, página {c['page']}]\n{c['content']}" for c in chunks)
    media = sum(c["similarity"] for c in chunks) / len(chunks) if chunks else 0.0
    return contexto, media


def contexto_para_texto(evidencias: list[dict]) -> str:
    return "\n\n".join(
        f"[Fonte: {e.get('fonte', '?')}, página {e.get('pagina', '?')}]\n{e.get('conteudo', '')}"
        for e in evidencias
    )


def _slug_arquivo(texto: str) -> str:
    normalizado = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    normalizado = re.sub(r"[^a-zA-Z0-9]+", "-", normalizado).strip("-").lower()
    return normalizado[:60] or "apostila"


def marca_do_nome(nome: str) -> str:
    """Marca curta e estável do nome original (FNV-1a de 32 bits, em hex).

    Serve para desempatar nomes que somem por inteiro na limpeza ASCII. Não é
    hash criptográfico: só precisa ser determinístico e idêntico ao do
    navegador (`marcaDoNome` em lib/texto.ts), porque a aba PDFs mostra o nome
    final antes do envio.
    """
    h = 0x811C9DC5
    for b in nome.encode("utf-8"):
        h = ((h ^ b) * 0x01000193) & 0xFFFFFFFF
    return format(h, "08x")


def nome_para_storage(nome: str) -> str:
    """Nome de arquivo que o Storage do Supabase aceita.

    O Storage recusa com 400 "InvalidKey" nomes com acento ou símbolo fora do
    ASCII ("Infográfico_RAG_final.pdf"), e o erro chegava ao admin como se fosse
    problema de chave ou de RLS. Aqui o acento sai ("á" → "a") e o que sobrar
    fora de letras, números, espaço e `._-()` vira "_". O mesmo nome é usado no
    Storage e como fonte no índice, então as citações mostram o nome limpo.
    """
    base = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode("ascii")
    base = re.sub(r"[^A-Za-z0-9 ._()\-]+", "_", base)
    base = re.sub(r"_{2,}", "_", base).strip(" _")
    # Nome inteiramente fora do ASCII ("日本語.pdf") não sobra nada da limpeza.
    # Chamar todos de "documento.pdf" fazia dois arquivos diferentes colidirem:
    # o segundo entrava como "versão atualizada do mesmo arquivo" e substituía o
    # primeiro. A marca do nome original desempata.
    raiz, ponto, extensao = base.rpartition(".")
    if ponto and not raiz.strip(" ._"):
        return f"documento-{marca_do_nome(nome)}.{extensao}"
    if not base:
        return f"documento-{marca_do_nome(nome)}"
    return base


# ── Markdown da apostila → blocos ───────────────────────────────────────────
#
# O modelo escreve a apostila em Markdown. Os dois formatos de saída (PDF e DOCX)
# leem a mesma lista de blocos, para que um recurso novo (caixa de destaque,
# tabela) não funcione num formato e saia cru no outro.

_RE_NUMERO_TITULO = re.compile(r"^\d+(\.\d+)*[.)]?\s+")
_RE_ITEM_NUMERADO = re.compile(r"^\d+[.)]\s+")
_RE_SEPARADOR_TABELA = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
_RE_ROTULO = re.compile(r"^\*\*([^*]{2,30}?)\s*:?\*\*\s*:?\s*")


def _celulas(linha: str) -> list[str]:
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def _blocos_markdown(texto: str) -> list[tuple[str, object]]:
    """Divide o Markdown em blocos: h1/h2/h3, p, ul, ol, quote, code, table.

    - quote: (rótulo, texto) — `> **Exemplo:** ...` vira uma caixa de destaque;
    - code: texto do bloco entre crases triplas (prompts de exemplo);
    - table: lista de linhas, a primeira é o cabeçalho.
    """
    blocos: list[tuple[str, object]] = []
    linhas = texto.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(linhas):
        bruta = linhas[i]
        linha = bruta.strip()
        if not linha or re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", linha):
            i += 1
            continue
        if linha.startswith("```"):
            corpo = []
            i += 1
            while i < len(linhas) and not linhas[i].strip().startswith("```"):
                corpo.append(linhas[i].rstrip())
                i += 1
            i += 1  # fecha a cerca (ou fim do texto)
            blocos.append(("code", "\n".join(corpo).strip("\n")))
            continue
        if linha.startswith("#"):
            nivel = len(linha) - len(linha.lstrip("#"))
            titulo = linha[nivel:].strip()
            if titulo:
                blocos.append((f"h{min(max(nivel, 1), 3)}", titulo))
            i += 1
            continue
        if linha.startswith("|"):
            linhas_tabela = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                if not _RE_SEPARADOR_TABELA.match(linhas[i].strip()):
                    linhas_tabela.append(_celulas(linhas[i]))
                i += 1
            if linhas_tabela:
                blocos.append(("table", linhas_tabela))
            continue
        if linha.startswith(">"):
            partes = []
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                partes.append(linhas[i].strip().lstrip(">").strip())
                i += 1
            corpo = " ".join(p for p in partes if p)
            m = _RE_ROTULO.match(corpo)
            rotulo, corpo = (m.group(1).strip(), corpo[m.end():]) if m else ("", corpo)
            blocos.append(("quote", (rotulo, corpo.strip())))
            continue
        if linha.startswith(("- ", "* ", "• ")):
            itens = []
            while i < len(linhas) and linhas[i].strip().startswith(("- ", "* ", "• ")):
                itens.append(linhas[i].strip()[2:].strip())
                i += 1
            blocos.append(("ul", itens))
            continue
        if _RE_ITEM_NUMERADO.match(linha):
            itens = []
            while i < len(linhas) and _RE_ITEM_NUMERADO.match(linhas[i].strip()):
                itens.append(_RE_ITEM_NUMERADO.sub("", linhas[i].strip(), count=1))
                i += 1
            blocos.append(("ol", itens))
            continue
        # Parágrafo: junta linhas seguidas até uma linha vazia ou outro bloco.
        partes = [linha]
        i += 1
        while i < len(linhas):
            prox = linhas[i].strip()
            if not prox or prox.startswith(("#", "```", "|", ">", "- ", "* ", "• ")) or _RE_ITEM_NUMERADO.match(prox):
                break
            partes.append(prox)
            i += 1
        blocos.append(("p", " ".join(partes)))
    return blocos


def _texto_puro(texto: str) -> str:
    """Tira as marcas de Markdown em linha (para títulos, sumário e DOCX)."""
    texto = texto.replace("**", "").replace("__", "").replace("`", "")
    return re.sub(r"(?<![\w*])\*(?!\s)([^*]+?)\*(?![\w*])", r"\1", texto)


def _texto_negrito(texto: str) -> str:
    """Prepara o texto para o `markdown=True` do fpdf2: mantém só o **negrito**.

    O fpdf2 também lê `__` (itálico) e `--` (sublinhado), que aparecem em texto
    comum ("--flag", "nome__campo") e bagunçariam a página.
    """
    texto = texto.replace("__", "").replace("`", "").replace("--", "–")
    texto = re.sub(r"(?<![\w*])\*(?!\s)([^*]+?)\*(?![\w*])", r"\1", texto)
    if texto.count("**") % 2:  # negrito sem fechamento: sai tudo normal
        texto = texto.replace("**", "")
    return texto


def _titulo_da_apostila(blocos: list[tuple[str, object]], tema: str) -> tuple[str, list[tuple[str, object]]]:
    """Usa o primeiro `# ` como título da capa; os demais `# ` viram seções."""
    titulo = tema
    resto = []
    for tipo, dado in blocos:
        if tipo == "h1" and titulo == tema and not resto:
            titulo = _texto_puro(str(dado))
            continue
        resto.append(("h2", dado) if tipo == "h1" else (tipo, dado))
    return titulo, resto


def _markdown_para_docx(texto: str, titulo: str) -> bytes:
    titulo, blocos = _titulo_da_apostila(_blocos_markdown(texto), titulo)
    doc = DocxDocument()
    doc.add_heading(titulo, level=0)
    for tipo, dado in blocos:
        if tipo in ("h2", "h3"):
            doc.add_heading(_RE_NUMERO_TITULO.sub("", _texto_puro(str(dado))), level=int(tipo[1]) - 1)
        elif tipo == "ul":
            for item in dado:  # type: ignore[union-attr]
                doc.add_paragraph(_texto_puro(item), style="List Bullet")
        elif tipo == "ol":
            for item in dado:  # type: ignore[union-attr]
                doc.add_paragraph(_texto_puro(item), style="List Number")
        elif tipo == "quote":
            rotulo, corpo = dado  # type: ignore[misc]
            p = doc.add_paragraph(style="Intense Quote")
            if rotulo:
                p.add_run(f"{rotulo}: ").bold = True
            p.add_run(_texto_puro(corpo))
        elif tipo == "code":
            p = doc.add_paragraph(str(dado))
            for run in p.runs:
                run.font.name = "Consolas"
        elif tipo == "table":
            linhas = dado  # type: ignore[assignment]
            colunas = max(len(r) for r in linhas)
            tabela = doc.add_table(rows=0, cols=colunas)
            tabela.style = "Light Grid Accent 1"
            for n, r in enumerate(linhas):
                celulas = tabela.add_row().cells
                for c in range(colunas):
                    celulas[c].text = _texto_puro(r[c]) if c < len(r) else ""
                    if n == 0:
                        for run in celulas[c].paragraphs[0].runs:
                            run.bold = True
        else:
            doc.add_paragraph(_texto_puro(str(dado)))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Fonte Unicode para o PDF. A Helvetica embutida no fpdf2 só aceita Latin-1, e
# o PDF quebrava (erro 500, depois de pagar a chamada de IA) com travessão,
# aspas curvas e "•", que os modelos usam o tempo todo (análise de 17/09/2026).
PASTA_FONTES = Path(__file__).resolve().parent / "fonts"
FONTE_PDF = PASTA_FONTES / "DejaVuSans.ttf"
FONTE_PDF_NEGRITO = PASTA_FONTES / "DejaVuSans-Bold.ttf"

# Plano B, sem a fonte: equivalentes em Latin-1 para os caracteres mais comuns.
# O que sobrar fora do Latin-1 (emoji, por exemplo) vira "?" em vez de erro.
_TROCAS_LATIN1 = str.maketrans({
    "—": "-", "–": "-", "‒": "-", "‐": "-", "‑": "-", "−": "-",
    "“": '"', "”": '"', "„": '"', "‘": "'", "’": "'", "‚": "'",
    "•": "-", "●": "-", "▪": "-", "…": "...",
    "→": "->", "←": "<-", "⇒": "=>", "↔": "<->",
    "≤": "<=", "≥": ">=", "≠": "!=", "≈": "~",
    "✓": "v", "✔": "v", "✗": "x", "✘": "x",
    "­": "", "​": "", "‌": "", "‍": "", "﻿": "",
})


def _para_latin1(texto: str) -> str:
    return texto.translate(_TROCAS_LATIN1).encode("latin-1", "replace").decode("latin-1")


# Paleta do PDF: sóbria, para impressão (o tema escuro do app não serve em papel).
_COR_PRIMARIA = (31, 58, 95)  # azul-marinho: capa, títulos
_COR_DESTAQUE = (42, 157, 143)  # verde-azulado: filetes, subtítulos
_COR_TEXTO = (40, 44, 52)
_COR_SUAVE = (120, 126, 138)

# Caixas de destaque: rótulo (sem acento, minúsculo) → (cor da barra, fundo).
_CAIXAS = [
    (("defini", "conceito"), (31, 58, 95), (234, 240, 248)),
    (("exemplo", "na pratica", "caso"), (42, 157, 143), (230, 245, 242)),
    (("atenc", "cuidado", "erro", "armadilha", "aviso"), (201, 125, 0), (253, 244, 225)),
    (("analogia",), (123, 76, 194), (240, 234, 250)),
    (("dica", "resumo", "importante", "lembre", "chave"), (46, 125, 50), (232, 245, 233)),
]
_CAIXA_PADRAO = ((90, 98, 110), (242, 243, 245))


def _cores_da_caixa(rotulo: str) -> tuple[tuple[int, int, int], tuple[int, int, int]]:
    chave = unicodedata.normalize("NFKD", rotulo).encode("ascii", "ignore").decode().lower()
    for palavras, cor, fundo in _CAIXAS:
        if any(p in chave for p in palavras):
            return cor, fundo
    return _CAIXA_PADRAO


def _data_hoje() -> str:
    from datetime import datetime

    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y")
    except Exception:
        return datetime.now().strftime("%d/%m/%Y")


def _markdown_para_pdf(texto: str, titulo: str) -> bytes:
    """Apostila em PDF com cara de material didático.

    Capa, sumário com páginas, seções numeradas, texto justificado, caixas de
    destaque (Definição, Exemplo, Atenção, Analogia, Dica), blocos de prompt,
    tabelas e rodapé com número de página.
    """
    from fpdf import FPDF  # import local: só carregado quando o formato pedido é pdf
    from fpdf.fonts import FontFace

    titulo, blocos = _titulo_da_apostila(_blocos_markdown(texto), titulo)
    titulo = _texto_puro(titulo)[:140] or "Apostila"

    class Apostila(FPDF):
        def header(self):
            if self.page_no() == 1:
                return
            self.set_font(familia, "", 8)
            self.set_text_color(*_COR_SUAVE)
            self.cell(0, 6, preparar(titulo[:90]), new_x="LMARGIN", new_y="NEXT")
            self.set_draw_color(220, 223, 228)
            self.set_line_width(0.2)
            self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
            self.ln(6)

        def footer(self):
            if self.page_no() == 1:
                return
            self.set_y(-15)
            self.set_font(familia, "", 8)
            self.set_text_color(*_COR_SUAVE)
            self.cell(0, 6, preparar("Engenharia de Prompt · apostila gerada com IA"), align="L")
            self.set_x(self.l_margin)
            self.cell(0, 6, f"{self.page_no()}", align="R")

    pdf = Apostila(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=22)

    familia, preparar = "Helvetica", _para_latin1
    if FONTE_PDF.exists() and FONTE_PDF_NEGRITO.exists():
        try:
            pdf.add_font("DejaVu", "", str(FONTE_PDF))
            pdf.add_font("DejaVu", "B", str(FONTE_PDF_NEGRITO))
            familia, preparar = "DejaVu", str
        except Exception:  # fonte corrompida: segue com o plano B em vez de falhar
            pass

    # new_x="LMARGIN" em todo multi_cell: o padrão do fpdf2 deixa o cursor na
    # borda direita, e a chamada seguinte com largura 0 quebra por falta de espaço.
    corpo_tam, corpo_alt = 10.5, 5.8

    def fonte(tamanho: float, negrito: bool = False, cor=_COR_TEXTO):
        pdf.set_font(familia, "B" if negrito else "", tamanho)
        pdf.set_text_color(*cor)

    # ── Capa ──
    pdf.add_page()
    pdf.set_fill_color(*_COR_PRIMARIA)
    pdf.rect(0, 0, pdf.w, 125, "F")
    pdf.set_fill_color(*_COR_DESTAQUE)
    pdf.rect(0, 125, pdf.w, 3, "F")
    pdf.set_xy(20, 40)
    fonte(10, True, (160, 225, 215))
    pdf.cell(0, 6, preparar("APOSTILA · ENGENHARIA DE PROMPT"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    fonte(26, True, (255, 255, 255))
    pdf.multi_cell(0, 11.5, preparar(titulo), align="L", new_x="LMARGIN", new_y="NEXT")
    pdf.set_y(142)
    fonte(10.5, False, _COR_TEXTO)
    pdf.multi_cell(
        0, 6,
        preparar(
            "Material didático gerado com inteligência artificial a partir do acervo do curso. "
            "As citações no formato [Fonte: arquivo, página N] indicam de onde veio cada informação — "
            "confira no material original antes de usar em provas ou trabalhos."
        ),
        new_x="LMARGIN", new_y="NEXT",
    )
    pdf.ln(4)
    fonte(9.5, False, _COR_SUAVE)
    pdf.cell(0, 6, preparar(f"Gerada em {_data_hoje()}"), new_x="LMARGIN", new_y="NEXT")

    # ── Sumário (preenchido pelo fpdf2 no fim, com as páginas reais) ──
    n_secoes = sum(1 for t, _ in blocos if t in ("h2", "h3"))

    def desenhar_sumario(pdf_, secoes):
        pdf_.set_x(pdf_.l_margin)
        fonte(18, True, _COR_PRIMARIA)
        pdf_.cell(0, 10, preparar("Sumário"), new_x="LMARGIN", new_y="NEXT")
        pdf_.set_draw_color(*_COR_DESTAQUE)
        pdf_.set_line_width(0.8)
        pdf_.line(pdf_.l_margin, pdf_.get_y() + 1, pdf_.l_margin + 25, pdf_.get_y() + 1)
        pdf_.ln(7)
        for s in secoes:
            recuo = 8 * s.level
            fonte(10.5 if s.level == 0 else 9.5, s.level == 0, _COR_TEXTO if s.level == 0 else _COR_SUAVE)
            pdf_.set_x(pdf_.l_margin + recuo)
            largura = pdf_.epw - recuo - 12
            nome = s.name if pdf_.get_string_width(s.name) <= largura else s.name[: int(len(s.name) * largura / pdf_.get_string_width(s.name)) - 1] + "…"
            pdf_.cell(largura, 7, preparar(nome))
            pdf_.cell(12, 7, str(s.page_number), align="R", new_x="LMARGIN", new_y="NEXT")

    pdf.add_page()
    if n_secoes:  # o fpdf2 já quebra a página depois do espaço reservado
        pdf.insert_toc_placeholder(desenhar_sumario, pages=max(1, -(-n_secoes // 30)))

    def caixa(rotulo: str, corpo: str, cor, fundo, codigo: bool = False):
        pad, barra = 4, 1.6
        largura = pdf.epw - 2 * pad - barra
        tam, alt = (9, 4.8) if codigo else (corpo_tam, corpo_alt)
        texto_caixa = preparar(corpo if codigo else _texto_negrito(corpo))
        fonte(tam)
        h_texto = pdf.multi_cell(
            largura, alt, texto_caixa, markdown=not codigo, align="L", dry_run=True, output="HEIGHT"
        )
        h_rotulo = 6.5 if rotulo else 0
        total = 2 * pad + h_rotulo + h_texto
        cabe_numa_pagina = total <= pdf.page_break_trigger - pdf.t_margin - 10
        if cabe_numa_pagina and pdf.will_page_break(total):
            pdf.add_page()
        x, y = pdf.l_margin, pdf.get_y()
        if cabe_numa_pagina:  # caixa maior que a página: sai sem fundo, só o texto
            pdf.set_fill_color(*fundo)
            pdf.rect(x, y, pdf.epw, total, "F")
            pdf.set_fill_color(*cor)
            pdf.rect(x, y, barra, total, "F")
        pdf.set_xy(x + barra + pad, y + pad)
        if rotulo:
            fonte(9, True, cor)
            pdf.cell(largura, h_rotulo, preparar(rotulo.upper()), new_x="LEFT", new_y="NEXT")
        fonte(tam, False, _COR_TEXTO)
        pdf.set_x(x + barra + pad)
        pdf.multi_cell(
            largura, alt, texto_caixa, markdown=not codigo, align="L", new_x="LMARGIN", new_y="NEXT"
        )
        if cabe_numa_pagina:
            pdf.set_y(y + total)
        pdf.ln(4)

    def lista(itens: list[str], numerada: bool):
        for n, item in enumerate(itens, 1):
            if pdf.will_page_break(corpo_alt):
                pdf.add_page()
            fonte(corpo_tam, numerada, _COR_DESTAQUE)
            pdf.set_x(pdf.l_margin + 3)
            pdf.cell(7, corpo_alt, f"{n}." if numerada else preparar("•"))
            fonte(corpo_tam)
            pdf.multi_cell(
                pdf.epw - 10, corpo_alt, preparar(_texto_negrito(item)),
                markdown=True, align="L", new_x="LMARGIN", new_y="NEXT",
            )
            pdf.ln(1)
        pdf.ln(2)

    def tabela(linhas: list[list[str]]):
        colunas = max(len(r) for r in linhas)
        linhas = [[preparar(_texto_puro(c)) for c in r] + [""] * (colunas - len(r)) for r in linhas]
        fonte(9)
        # A tabela herda a cor de preenchimento e de linha do que veio antes
        # (a barra da última caixa pintava as linhas); fixa as duas aqui.
        pdf.set_fill_color(245, 247, 250)
        pdf.set_draw_color(215, 219, 225)
        pdf.set_line_width(0.2)
        try:
            with pdf.table(
                line_height=5.2,
                text_align="LEFT",
                borders_layout="HORIZONTAL_LINES",
                cell_fill_color=(245, 247, 250),
                cell_fill_mode="ROWS",
                headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=_COR_PRIMARIA),
            ) as t:
                for r in linhas:
                    linha = t.row()
                    for c in r:
                        linha.cell(c)
        except Exception:  # tabela torta demais para o fpdf2: sai como texto
            for r in linhas:
                pdf.multi_cell(0, corpo_alt, " | ".join(r), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)

    n2 = n3 = 0
    for tipo, dado in blocos:
        if tipo == "h2":
            n2, n3 = n2 + 1, 0
            nome = _RE_NUMERO_TITULO.sub("", _texto_puro(str(dado)))
            if pdf.get_y() > pdf.t_margin + 20 and pdf.will_page_break(35):
                pdf.add_page()
            elif pdf.get_y() > pdf.t_margin + 20:
                pdf.ln(5)
            pdf.start_section(preparar(f"{n2}. {nome}"), level=0)
            fonte(15, True, _COR_PRIMARIA)
            pdf.multi_cell(0, 8, preparar(f"{n2}. {nome}"), align="L", new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(*_COR_DESTAQUE)
            pdf.set_line_width(0.8)
            pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.l_margin + 25, pdf.get_y() + 1)
            pdf.ln(5)
        elif tipo == "h3":
            n3 += 1
            nome = _RE_NUMERO_TITULO.sub("", _texto_puro(str(dado)))
            rotulo = f"{n2}.{n3} {nome}" if n2 else nome
            if pdf.will_page_break(25):
                pdf.add_page()
            pdf.ln(2)
            pdf.start_section(preparar(rotulo), level=1)
            fonte(12, True, (32, 110, 100))
            pdf.multi_cell(0, 7, preparar(rotulo), align="L", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
        elif tipo == "p":
            fonte(corpo_tam)
            pdf.multi_cell(
                0, corpo_alt, preparar(_texto_negrito(str(dado))),
                markdown=True, align="J", new_x="LMARGIN", new_y="NEXT",
            )
            pdf.ln(2.5)
        elif tipo in ("ul", "ol"):
            lista(dado, tipo == "ol")  # type: ignore[arg-type]
        elif tipo == "quote":
            rotulo, corpo = dado  # type: ignore[misc]
            caixa(rotulo, corpo, *_cores_da_caixa(rotulo))
        elif tipo == "code":
            caixa("Prompt / exemplo", str(dado), (90, 98, 110), (246, 247, 249), codigo=True)
        elif tipo == "table":
            tabela(dado)  # type: ignore[arg-type]
    return bytes(pdf.output())


def _mimetype(nome: str) -> str:
    if nome.lower().endswith(".pdf"):
        return "application/pdf"
    if nome.lower().endswith(".md"):
        return "text/markdown; charset=utf-8"
    if nome.lower().endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if nome.lower().endswith(".pptx"):
        return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    return "text/plain; charset=utf-8"
