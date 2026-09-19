"""PDF do Gerador de Apostilas com caracteres fora do Latin-1 (análise de 17/09/2026, A1-04).

    .venv/Scripts/python.exe -m pytest tests/test_apostila_pdf.py -v
"""
import io
import sys
from pathlib import Path

import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "api"))

from pypdf import PdfReader  # noqa: E402

from _app import documentos  # noqa: E402

# Antes da correção, todos levantavam FPDFUnicodeEncodingException na Helvetica.
CASOS = [
    "Inovação — conceito",
    "Intervalo 2010–2020",
    "Aspas “curvas” e ‘simples’",
    "Lista • item",
    "Reticências…",
    "Seta → resultado",
    "Hífen­inseparável",
    "x ≥ 3 e y ≤ 5",
    "Emoji 🚀 no texto",
    "- item com **negrito** no meio",
]


def _gerar(texto: str, titulo: str = "Tema — IA") -> bytes:
    pdf = documentos._markdown_para_pdf(texto, titulo)
    assert pdf.startswith(b"%PDF")
    return pdf


@pytest.mark.parametrize("texto", CASOS)
def test_pdf_aceita_caracteres_comuns_de_ia(texto):
    _gerar(f"# Título\n\n{texto}\n\n## Seção\n\n{texto}")


@pytest.mark.parametrize("texto", CASOS)
def test_pdf_sem_fonte_unicode_cai_no_plano_b(monkeypatch, tmp_path, texto):
    monkeypatch.setattr(documentos, "FONTE_PDF", tmp_path / "nao-existe.ttf")
    _gerar(texto)


def test_pdf_remove_negrito_do_markdown():
    pdf = _gerar("Um **termo** importante")
    texto = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "**" not in texto
    assert "termo" in texto


def test_plano_b_troca_por_equivalentes():
    assert documentos._para_latin1("a — b “c” • d… →") == 'a - b "c" - d... ->'
    assert documentos._para_latin1("ação 🚀") == "ação ?"


def test_fonte_unicode_esta_no_pacote():
    """A fonte precisa estar em api/_app/fonts, que o vercel.json empacota (`_app/**`).
    Sem ela o PDF ainda sai, mas com símbolos trocados por equivalentes simples."""
    assert documentos.FONTE_PDF.exists(), f"faltando {documentos.FONTE_PDF}"
    assert documentos.FONTE_PDF_NEGRITO.exists(), f"faltando {documentos.FONTE_PDF_NEGRITO}"


APOSTILA = """# Título da apostila

## Objetivos de aprendizagem
- **Entender** o tema
- Aplicar na prática

## Conceito
Parágrafo com **negrito** e `código` que
continua na linha seguinte.

> **Definição:** texto da definição.

> **Analogia:** como no futebol.

```
Pense passo a passo.
```

| Termo | Significado |
|---|---|
| CoT | Chain-of-Thought |

### Subseção
1. primeiro
2. segundo
"""


def test_blocos_reconhecem_a_estrutura():
    blocos = documentos._blocos_markdown(APOSTILA)
    tipos = [t for t, _ in blocos]
    assert tipos == ["h1", "h2", "ul", "h2", "p", "quote", "quote", "code", "table", "h3", "ol"]
    assert blocos[4][1] == "Parágrafo com **negrito** e `código` que continua na linha seguinte."
    assert blocos[5][1] == ("Definição", "texto da definição.")
    assert blocos[7][1] == "Pense passo a passo."
    assert blocos[8][1] == [["Termo", "Significado"], ["CoT", "Chain-of-Thought"]]


def test_primeiro_h1_vira_titulo_da_capa():
    titulo, resto = documentos._titulo_da_apostila(documentos._blocos_markdown(APOSTILA), "tema")
    assert titulo == "Título da apostila"
    assert all(t != "h1" for t, _ in resto)


def test_pdf_tem_capa_sumario_e_secoes_numeradas():
    pdf = _gerar(APOSTILA)
    paginas = [p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages]
    assert "Título da apostila" in paginas[0]
    assert "Sumário" in paginas[1]
    corpo = "".join(paginas[2:])
    assert "1. Objetivos de aprendizagem" in corpo
    assert "2.1 Subseção" in corpo
    assert "DEFINIÇÃO" in corpo and "ANALOGIA" in corpo
    assert "**" not in corpo and "`" not in corpo


def test_pdf_longo_com_muitas_secoes_e_caixa_maior_que_a_pagina():
    secoes = "\n\n".join(f"## Seção {n}\n\n### Parte {n}\n\ntexto" for n in range(40))
    caixa = "> **Exemplo:** " + "palavra " * 1500
    _gerar(f"# T\n\n{secoes}\n\n{caixa}")


def test_pdf_sem_titulos_nao_tem_sumario():
    paginas = PdfReader(io.BytesIO(_gerar("Só um parágrafo."))).pages
    assert len(paginas) == 2  # capa + conteúdo


def test_docx_usa_os_mesmos_blocos():
    from docx import Document

    doc = Document(io.BytesIO(documentos._markdown_para_docx(APOSTILA, "tema")))
    textos = [p.text for p in doc.paragraphs]
    assert textos[0] == "Título da apostila"
    assert "Definição: texto da definição." in textos
    assert not any("**" in t or t.startswith(">") for t in textos)
    assert len(doc.tables) == 1


# O fpdf2 exige que as páginas reservadas ao sumário sejam exatamente as que ele
# ocupa. A estimativa antiga (30 linhas por página) errava para 31-32, 61-68,
# 91-104... seções, e a rota devolvia 500 depois de gastar a chamada de IA.
@pytest.mark.parametrize("n_secoes", [1, 29, 30, 31, 32, 33, 35, 36, 61, 65, 67, 68, 69, 100, 101, 104, 105, 130])
def test_pdf_sumario_fecha_para_qualquer_quantidade_de_secoes(n_secoes):
    md = "\n\n".join(f"## Seção {n}\n\ntexto" for n in range(n_secoes))
    pdf = _gerar(f"# T\n\n{md}")
    corpo = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert f"{n_secoes}. Seção {n_secoes - 1}" in corpo


# O modelo às vezes devolve caracteres de controle (o projeto já limpa NUL no
# texto extraído). O python-docx recusa esses caracteres e a rota dava 500.
@pytest.mark.parametrize("sujeira", ["\x00", "\x07", "\x1b", "\x0b", "\x0c"])
def test_docx_aceita_caracteres_de_controle(sujeira):
    from docx import Document

    md = f"# Tema{sujeira}\n\nTexto{sujeira}com{sujeira}lixo.\n\n- item{sujeira}\n\n| a | b |\n|---|---|\n| x{sujeira} | y |"
    doc = Document(io.BytesIO(documentos._markdown_para_docx(md, f"Tema{sujeira}")))
    assert any("Textocomlixo." in p.text for p in doc.paragraphs)
